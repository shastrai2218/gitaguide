# -*- coding: utf-8 -*-
"""Log Book for Capstone Project (NMIMS template): cover page + 15 weekly entries, two per page."""
import os
import datetime as dt

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = ("GitaGuide: An Emotion-Aware, Explainable and Safety-Gated Guidance System "
         "Grounded in the Bhagavad Gita and the Arthashastra")
START = dt.date(2026, 7, 13)  # Monday of week 1; week 7 starts on 24 Aug 2026, the first commit

WEEKS = [
    ["Chose the project topic: a guidance system built on the Bhagavad Gita and the Arthashastra.",
     "Read news reports about Gita chatbots that gave harmful answers, to understand the problem.",
     "Wrote down what the system must do: check safety first, understand emotion, and cite real verses.",
     "Prepared a short summary of the idea to show my guide."],
    ["Started reading research papers on scripture-based chatbots and on AI for emotional support.",
     "Read how language models are used in mental health and what risks the reviews report.",
     "Made a table of the papers I read, with what each one does and what it does not do.",
     "Noticed early that I could not find a system that combined emotion, two scriptures and a safety check."],
    ["Read more papers on knowledge graphs, graph-based retrieval and crisis detection.",
     "Read papers on bias, emotional dependency and the harm caused by sudden refusals in chatbots.",
     "Kept adding to my comparison table of existing systems.",
     "Started writing the literature review."],
    ["Searched again on purpose for papers that disagree with my idea, such as work on bias and dependency.",
     "Found five more papers that changed some of my statements.",
     "Finished reviewing 34 papers and wrote down the research gap.",
     "Wrote the five requirements that my system has to meet."],
    ["Designed the order of steps for one message: emotion, intent, safety check, memory, retrieval, answer, output check.",
     "Decided that the safety check must run before everything else.",
     "Chose the tools: FastAPI, a Neo4j graph, PostgreSQL, a hosted language model and a React interface.",
     "Drew the first block diagram of the system."],
    ["Set up the project folders, the code layout and the testing setup.",
     "Looked for free and legal sources for the text of the Gita and the Arthashastra.",
     "Planned how to label verses by emotion and theme and how to connect them in a graph.",
     "Planned how I would test the system: safety prompts and a set of test queries with correct verses."],
    ["Made the first working version that answers with cited verses from start to end (24 Aug).",
     "Added the Arthashastra (3,135 sutras) and rebuilt the knowledge graph with both texts.",
     "Left out very common labels when connecting verses, so the graph stays small enough for the free tier.",
     "Changed the Gita translation to Purohit’s, because it is public domain.",
     "Wrote 40 test queries with the correct verses marked, and recorded the first result."],
    ["Checked the live system with the database and the graph, and wrote down what I found (2 Sep).",
     "Made the crisis check stronger and added a backup language model (5 Sep).",
     "Built the first chat screen with a consent page and a crisis panel.",
     "Made the system run without PyTorch and measured how long each step takes.",
     "Added rate limits, consent saved on the server, feedback buttons and a time limit for answers.",
     "Prepared the deployment files, a smoke test and the first notes for testers (6 Sep)."],
    ["Made the embedding model smaller (INT8) so the system fits in 512 MiB of memory; the peak was 429 MiB.",
     "Made the system handle one question at a time to keep memory low.",
     "Added a check that catches broken answers, such as repeated words.",
     "Made replies match the user’s language: English, Hindi or Hinglish.",
     "Changed answers to be short and conversational, with the verse reference shown beside each verse.",
     "Added a router so that a greeting does not get a verse (13 Sep)."],
    ["Made the router faster by handling common phrases without calling the model.",
     "Redesigned the screen with a sidebar, regenerate and edit buttons, and optional sign-in.",
     "Fixed a deployment bug where the data folder was missing from the image.",
     "Moved the API to an AWS EC2 server in Mumbai; one turn became about 12 seconds faster.",
     "Added a keep-alive so the free graph database does not pause.",
     "Fixed a bug where chats stayed visible after signing out (20 Sep)."],
    ["Went through all the measurements so far and listed the weak points: retrieval, intent and faithfulness.",
     "Decided which extra tests I needed before writing the report.",
     "Planned a larger test set, a human rating study and a safety red-team set.",
     "Cleaned up the code and my notes."],
    ["Wrote the tools to compare retrieval variants fairly: keyword only, vector only and with the graph.",
     "Wrote a guide and tools so that two people can judge the retrieved verses.",
     "Wrote 33 red-team prompts, including harmful questions framed as duty, and ran them on the system.",
     "Ran an automatic faithfulness check on 69 answers; one judge model failed, so I used the one that worked.",
     "Wrote my literature review paper and a cover letter (30 Sep)."],
    ["Judged the retrieved verses for 100 queries together with a second person.",
     "Rated 30 answers with the second person on relevance, faithfulness, helpfulness and appropriateness.",
     "Started writing the research paper.",
     "Wrote the first draft of the interim report."],
    ["Re-ran the retrieval tests, turning one part off at a time, and compared the results with a statistical test.",
     "Tested the emotion model on 1,500 public GoEmotions examples (macro-F1 0.451).",
     "Sent 100 test queries to the live system to check intent classification (65 of 100 correct).",
     "Re-ran the 33 red-team prompts on the live system and found 2 answers that wrongly allowed violence.",
     "Added a fifth output rule to block such answers and 4 new tests (tests passing went from 907 to 911).",
     "Finished the interim report (17 Oct)."],
    ["Revised the research paper: fixed tables, numbers and wording after review comments.",
     "Made the emotion table list all 28 labels and saved the full results in a file.",
     "Wrote the Black Book following the college template (55 pages).",
     "Made the language of the reports simpler and plainer.",
     "Prepared this log book."],
]
assert len(WEEKS) == 15


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), color)
    tcPr.append(s)


def set_borders(table, sz=6, color="000000"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def cell_text(cell, text, size=11, bold=False, align=None, highlight=False, min_h=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if highlight:
            r.font.highlight_color = 7
    tr = cell._tc.getparent()
    if min_h:
        trPr = tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(int(min_h * 567)))
        h.set(qn("w:hRule"), "atLeast")
        trPr.append(h)


def cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def P(D, text="", size=12, bold=False, align=None, after=6, before=0, highlight=False, underline=False):
    p = D.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.0
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.underline = underline
        if highlight:
            r.font.highlight_color = 7
    return p


D = docx.Document()
sec = D.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # US Letter, as in the template
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.4)
st = D.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")

# ---------------------------------------------------------------- cover page
p = D.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
p.add_run().add_picture(os.path.join(HERE, "nmims_logo.jpeg"), width=Cm(4.2))
P(D, "Mukesh Patel School of Technology Management & Engineering", 16, True, "c", after=22)
P(D, "Log Book for Capstone Project", 15, True, "c", after=26)
for label, val, hl in (("Department : ", "Computer Engineering", False),
                       ("Program : ", "B.Tech in Computer Engineering", False),
                       ("Semester : ", "[Semester]", True)):
    p = P(D, "", after=2)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(14)
    r2 = p.add_run(val)
    r2.font.size = Pt(14)
    r2.bold = True
    if hl:
        r2.font.highlight_color = 7
p = P(D, "", after=16)
r = p.add_run("TITLE OF THE PROJECT : ")
r.bold = True
r.font.size = Pt(14)
r = p.add_run(TITLE)
r.bold = True
r.font.size = Pt(13)
p = P(D, "", after=18)
p.add_run("Name of the Faculty Mentor: ").font.size = Pt(12)
r = p.add_run("Dr. Suresh Kurumbanshi")
r.bold = True
r.font.size = Pt(12)
r.underline = True

P(D, "STUDENT DETAILS", 13, True, "c", after=8)
t = D.add_table(rows=2, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
set_borders(t)
widths = (2.6, 6.8, 3.6, 4.5)
for c, h, w in zip(t.rows[0].cells, ("", "NAME", "ROLL NO.", "CONTACT"), widths):
    cell_text(c, h, 11, True, "c", min_h=0.8)
    c.width = Cm(w)
for c, v, w, hl in zip(t.rows[1].cells, ("Student 1", "Harikrishna Panchal", "[Roll No.]",
                                         "Harikrishna.Panchal505@nmims.in"), widths, (False, False, True, False)):
    cell_text(c, v, 10.5 if "@" in v else 11, v == "Student 1", highlight=hl, min_h=1.2)
    c.width = Cm(w)
P(D, "", after=10)
P(D, "Name of group members:", 11, after=6)
t = D.add_table(rows=2, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
set_borders(t, 8, "4472C4")
widths = (2.6, 6.8, 3.6, 4.5)
for c, h, w in zip(t.rows[0].cells, ("Roll No.", "Name", "Sign (Students)", "Grades (By Faculty Mentor)"), widths):
    cell_text(c, h, 11, True, "c", min_h=0.8)
    shade(c, "D6E3F0")
    c.width = Cm(w)
for c, v, w, hl in zip(t.rows[1].cells, ("[Roll No.]", "Harikrishna Panchal", "", ""), widths, (True, False, False, False)):
    cell_text(c, v, 11, highlight=hl, min_h=1.0)
    c.width = Cm(w)

# ---------------------------------------------------------------- weekly pages
for i, work in enumerate(WEEKS):
    if i % 2 == 0:
        D.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        D.paragraphs[-1].paragraph_format.space_after = Pt(0)
        D.paragraphs[-1].paragraph_format.line_spacing = 0.1
    monday = START + dt.timedelta(weeks=i)
    sunday = monday + dt.timedelta(days=6)
    rows = [("Week No: %d" % (i + 1), 0.68), ("Date of Reporting:", 0.68),
            ("From %s to %s" % (monday.strftime("%d/%m/%Y"), sunday.strftime("%d/%m/%Y")), 0.68),
            ("Work carried out", 0.68)]
    n_work = 6
    for k in range(n_work):
        rows.append((("• " + work[k]) if k < len(work) else "", 0.68))
    rows += [("Faculty Remark", 0.68), ("", 0.68), ("", 0.68), ("", 0.68),
             ("Faculty Signature with date", 0.68), ("", 0.68)]
    t = D.add_table(rows=len(rows), cols=1)
    t.autofit = False
    set_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (txt, h) in zip(t.rows, rows):
        cell_text(row.cells[0], txt, 10.5, False, min_h=h)
        row.cells[0].width = Cm(17.5)
        cant_split(row)
    if i % 2 == 0 and i != len(WEEKS) - 1:
        P(D, "", after=0, size=6)

D.save(os.path.join(HERE, "LogBook_GitaGuide.docx"))
print("saved")
