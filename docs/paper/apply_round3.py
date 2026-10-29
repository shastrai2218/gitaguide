# -*- coding: utf-8 -*-
"""Third revision pass, addressing reviewer feedback on the round-2 paper:
the abstract, Table III, Table XIII, and two discussion paragraphs still
described the red-team/ablation work as not-yet-done after it was done.
Operates directly on the current Research_Paper.docx (already carries
rounds 1-2); does not re-run apply_main.py.
"""
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\Projects\GitaGuide\gitaguide\docs\paper\Research_Paper.docx"
D = docx.Document(SRC)


def find(substr, start=0):
    for i in range(start, len(D.paragraphs)):
        if substr in D.paragraphs[i].text:
            return i
    raise ValueError(f"not found: {substr!r}")


def replace_text(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return
    raise AssertionError(f"{old!r} not found within any single run of: {p.text!r}")


# ============================================================ 1. Abstract: fix the contradictory red-team sentence + explicit deployment caveat
i = find("A red-team test of 25 dharma-framed violence prompts found none judged as endorsing harm, though one of two planned judges failed outright.")
replace_text(
    D.paragraphs[i],
    "A red-team test of 25 dharma-framed violence prompts found none judged as endorsing harm, "
    "though one of two planned judges failed outright.",
    "A red-team test of 25 dharma-framed violence prompts was first graded by one working judge, "
    "which found no clean endorsement of harm; a second, working judge we added later (Section V-G) "
    "disagreed on two of the 25, finding the system had operationalised scripture-sanctioned "
    "violence rather than refusing it.",
)
i = find("and a second, working red-team judge found two real safety failures the first judge missed")
replace_text(
    D.paragraphs[i],
    "and a second, working red-team judge found two real safety failures the first judge missed "
    "\u2014 both root-caused and fixed, with regression tests, within this revision.",
    "and that second judge's two findings were root-caused to a gap in the output gate and fixed "
    "locally, with regression tests confirming the fix against the same captured replies; the fix "
    "has not been deployed to the production instance this paper otherwise measures.",
)

# ============================================================ 2. Table III: add the fifth output-gate rule with explicit deployment status
for t in D.tables:
    hdr = [c.text for c in t.rows[0].cells]
    if hdr == ["Layer", "Mechanism", "Effect"]:
        for row in t.rows:
            if row.cells[0].text == "Output":
                mech_cell, eff_cell = row.cells[1], row.cells[2]
                for cell, new_text in (
                    (mech_cell, "5 patterns on the generated answer: harm methods, discouraging "
                                 "treatment, moralising, dismissal, and (added after a second-judge "
                                 "red-team pass, Section V-G) operationalising a historical violence "
                                 "or torture method instead of refusing it"),
                    (eff_cell, "Replace answer. The 5th pattern is committed and unit-tested "
                                "(911 tests) but NOT deployed to the production instance measured "
                                "elsewhere in this paper"),
                ):
                    for para in cell.paragraphs:
                        for rn in list(para.runs):
                            rn.text = ""
                    cell.paragraphs[0].runs[0].text = new_text
                    cell.paragraphs[0].runs[0].font.size = Pt(8)
                print("Table III output row updated")
        break

# ============================================================ 3. Table XIII: update the two stale claim rows
for t in D.tables:
    hdr = [c.text for c in t.rows[0].cells]
    if hdr == ["Claim", "Evidence", "Limit"]:
        for row in t.rows:
            claim = row.cells[0].text
            if claim.startswith("Individual components"):
                new_evidence = (
                    "Measured (Table XI): BM25-only is significantly worse than the full "
                    "pipeline (p < 0.001, n = 81); a graph-expansion variant is significantly "
                    "worse, not better (p = 0.001); other variants are not yet distinguishable "
                    "from the full pipeline at this sample size"
                )
                new_limit = (
                    "Confirms the dense encoder/fusion matter and that naive graph expansion "
                    "hurts; does not yet isolate the gate or emotion-aware routing alone"
                )
                for cell, text in ((row.cells[1], new_evidence), (row.cells[2], new_limit)):
                    for para in cell.paragraphs:
                        for rn in list(para.runs):
                            rn.text = ""
                    cell.paragraphs[0].runs[0].text = text
                    cell.paragraphs[0].runs[0].font.size = Pt(8)
                print("updated 'individual components' row")
            elif claim.startswith("Intent and emotion classification"):
                new_evidence = (
                    "Measured: per-class P/R/F1 and a confusion matrix from a live, 100-query "
                    "run (Tables VII\u2013VIII); a held-out GoEmotions evaluation for the emotion "
                    "model, n = 1,500 (Table II)"
                )
                new_limit = (
                    "Overall accuracy is only 0.650 live (0.571 macro-F1 excluding two "
                    "zero-support classes); scripture_question is badly confused with "
                    "seek_guidance/professional_dilemma (F1 0.118); a 3% crisis-overtrigger "
                    "rate on benign queries; emotion intensity (equation 1) is still not "
                    "separately evaluated"
                )
                for cell, text in ((row.cells[1], new_evidence), (row.cells[2], new_limit)):
                    for para in cell.paragraphs:
                        for rn in list(para.runs):
                            rn.text = ""
                    cell.paragraphs[0].runs[0].text = text
                    cell.paragraphs[0].runs[0].font.size = Pt(8)
                print("updated 'intent and emotion' row")
        break

# ============================================================ 4. "What is better" discussion: fix the stale "not yet isolated by ablation" claim
i = find("These architectural choices have not yet been isolated by ablation")
replace_text(
    D.paragraphs[i],
    "These architectural choices have not yet been isolated by ablation: the numbers in Section V "
    "describe the full pipeline, not the marginal contribution of the knowledge graph, the safety "
    "gate, or emotion-aware routing taken alone (Section VI-C).",
    "Component ablations (Table XI) now isolate some of this: the dense encoder and fusion carry "
    "the measured retrieval accuracy, lexical search alone is significantly worse, and a naive "
    "graph-expansion variant is significantly worse, not better. The safety gate and emotion-aware "
    "routing are not yet isolated the same way (Section VI-C, item iv).",
)

# ============================================================ 5. "Comparison with the precedents": fix the stale single-judge claim
i = find("A first test of 25 dharma-framed prompts (Section V-G) found none that a single working judge")
replace_text(
    D.paragraphs[i],
    "A first test of 25 dharma-framed prompts (Section V-G) found none that a single working "
    "judge labelled as endorsing harm, though the second planned judge failed throughout, so "
    "this is not yet an independently confirmed result, and we do not claim to have solved it.",
    "A first test of 25 dharma-framed prompts (Section V-G) found none that a single working "
    "judge labelled as endorsing harm; a second, working judge we added later disagreed on two "
    "of the 25, finding the system had operationalised scripture-sanctioned violence instead of "
    "refusing it. Both cases are now fixed in the codebase and regression-tested against the "
    "same captured replies, but the fix is not deployed to production, so we still do not claim "
    "to have solved this for the live system.",
)

# ============================================================ 6. Strengthen the red-team methodology paragraph: manual assessment, the 20/3/2 split in prose, regression vs. live distinction
i = find("The second judge agrees with the first on the headline claim")
p = D.paragraphs[i]
for r in p.runs:
    if r.text.startswith("The second judge agrees"):
        r.text = (
            "This 33-prompt set was graded manually, by reading each live reply against the "
            "same three-way rubric used for the first judge (no endorsement / ambiguous / "
            "endorses harm), not by an automated scoring pipeline. Of the 25 harm-framed "
            "prompts, 20 were clear refusals, 3 were ambiguous (a soft redirect without an "
            "explicit refusal), and 2 were judged as endorsing harm outright (Table XII). "
            "The second judge agrees with the first on the headline claim"
        )
        break
else:
    raise AssertionError("anchor run not found")

i2 = find("Both failures are now fixed, not just reported.")
p2 = D.paragraphs[i2]
for r in p2.runs:
    if "Re-running it against all 33 real replies from this red-team pass flags exactly the two true failures" in r.text:
        r.text = r.text.replace(
            "Re-running it against all 33 real replies from this red-team pass flags exactly the "
            "two true failures and nothing else: zero false positives on the other 31, including "
            "the correct Hinglish refusal and all 8 controls.",
            "This is a regression replay, not a fresh live estimate: it re-runs the new rule "
            "against the same 33 responses already captured above, and flags exactly the two true "
            "failures and nothing else \u2014 zero false positives on the other 31, including the "
            "correct Hinglish refusal and all 8 controls. It confirms the fix works against the "
            "failures that produced it; it is not independent evidence of the false-positive or "
            "false-negative rate on new prompts, which future work item (iii) still calls for.",
        )
        break
else:
    raise AssertionError("regression-replay anchor run not found")

# ============================================================ 7. Conclusion: mention the fix status explicitly
i = find("We regard the negative results and the measurement pitfalls as findings in their own right")
p = D.paragraphs[i]
for r in p.runs:
    if "a measured system is not necessarily the deployed one." in r.text:
        r.text = r.text.replace(
            "a measured system is not necessarily the deployed one.",
            "a measured system is not necessarily the deployed one \u2014 a second-judge red-team "
            "pass found two real safety failures this revision fixed and regression-tested in the "
            "codebase, but has not yet deployed to the production instance the rest of this paper "
            "measures.",
        )
        break
else:
    raise AssertionError("conclusion anchor run not found")

# ============================================================ 8. Future work item (iii): reflect what is now partly done
i = find("Future work, in priority order, before any claim of validated accuracy:")
p = D.paragraphs[i]
for r in p.runs:
    if "(iii) safety \u2014 repeat the red team with an independently constructed prompt set" in r.text:
        r.text = r.text.replace(
            "(iii) safety \u2014 repeat the red team with an independently constructed prompt set "
            "(indirect, ambiguous, misspelled and multilingual phrasings), report false positives "
            "and negatives with confidence intervals, repeat fault injection, and add independent "
            "human review rather than a single LLM judge;",
            "(iii) safety \u2014 deploy the output-gate fix to production and re-measure; then repeat "
            "the red team with an independently constructed prompt set (indirect, ambiguous, "
            "misspelled and multilingual phrasings, not just a second judge on the same 33), "
            "report false positives and negatives with confidence intervals, and repeat fault "
            "injection;",
        )
        break
else:
    raise AssertionError("future-work anchor run not found")

D.save(SRC)
print("round 3 saved")
