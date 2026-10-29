# -*- coding: utf-8 -*-
import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import lib
from lib import (D, HERE, J, NOTE, P, TBL, UNNUM_TITLE, YEAR, bookmark, old_registry, page_of)

TITLE = ("GitaGuide: An Emotion-Aware, Explainable and Safety-Gated Guidance System "
         "Grounded in the Bhagavad Gita and the Arthashastra")


def logo(width):
    p = D.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(os.path.join(HERE, "img", "nmims_logo.jpeg"), width=Cm(width))
    return p


def cover():
    P("", after=0, line=1.0)
    P(TITLE, b=True, al="c", sz=24, after=22, line=1.15)
    P("Submitted in partial fulfillment of the requirement for the award of\nDegree of Bachelor of Technology in\nComputer Engineering",
      b=True, al="c", sz=14, after=16, line=1.15)
    P("Submitted To", b=True, al="c", sz=12, after=4, line=1.0)
    logo(4.3).paragraph_format.space_after = Pt(10)
    P("SVKM’s NMIMS,\nMukesh Patel School of Technology Management & Engineering,\nShirpur Campus (M.H.)",
      b=True, al="c", sz=12, after=14, line=1.15)
    P("Submitted By:", b=True, al="c", sz=12, after=2, line=1.0)
    p = P("Harikrishna Panchal    Enrollment No – ", b=True, al="c", sz=14, after=14, line=1.0)
    r = p.add_run("[Enrollment No.]")
    r.bold = True
    r.font.size = Pt(14)
    r.font.highlight_color = 7
    P("Under The Supervision Of:", b=True, al="c", sz=12, after=2, line=1.0)
    P("Dr. Suresh Kurumbanshi", b=True, al="c", sz=14, after=0, line=1.0)
    P("(Head of Department, Department of Computer Engineering)", b=True, al="c", sz=14, after=34, line=1.0)
    P("DEPARTMENT OF COMPUTER ENGINEERING", b=True, al="c", sz=14, after=2, line=1.0)
    P("Mukesh Patel School of Technology Management & Engineering, Shirpur", b=True, al="c", sz=12, after=2, line=1.0)
    P(f"ACADEMIC SESSION: {YEAR}", b=True, al="c", sz=12, after=0, line=1.0)


def certificate():
    p = UNNUM_TITLE("CERTIFICATE", "certificate", size=16, front=True, underline=True)
    logo(3.4)
    P("", after=2, line=1.0)
    p = P("", al="j")
    p.add_run("This is to certify that the project entitled ")
    r = p.add_run("“" + TITLE + "”")
    r.bold = True
    p.add_run(" has been done by")
    p = P("“Harikrishna Panchal” (Roll No.: ", b=True, al="c", sz=14)
    r = p.add_run("[Roll No.]")
    r.bold = True
    r.font.size = Pt(14)
    r.font.highlight_color = 7
    p.add_run(")").bold = True
    p.runs[-1].font.size = Pt(14)
    p = P("", al="j")
    p.add_run("under my guidance and supervision").bold = True
    p.add_run(" and has been submitted in partial fulfillment of the degree of “")
    p.add_run("Bachelor of Technology in Computer Engineering").bold = True
    p.add_run("” of MPSTME, SVKM’S NMIMS (Deemed-to-be university), Mumbai (M.H.), India.")
    P("", after=18, line=1.0)
    t = D.add_table(rows=1, cols=2)
    for c, (a, b2) in zip(t.rows[0].cells, (("_______________________", "Project Mentor (Internal Guide)\nDr. Suresh Kurumbanshi"),
                                            ("_______________________", "Examiner"))):
        c.text = ""
        pp = c.paragraphs[0]
        pp.paragraph_format.line_spacing = 1.0
        pp.add_run(a + "\n").bold = True
        for ln in b2.split("\n"):
            rr = pp.add_run(ln + "\n")
            rr.bold = True
            rr.font.size = Pt(12)
    P("", after=6, line=1.0)
    P("Date:", b=True, sz=12, after=2, line=1.0)
    P("Place: Shirpur", b=True, sz=12, after=14, line=1.0)
    P("_______________________", b=True, al="r", after=0, line=1.0)
    P("H.O.D.", b=True, al="r", sz=12, after=16, line=1.0)
    P("DEPARTMENT OF COMPUTER ENGINEERING", b=True, al="c", sz=14, after=2, line=1.0)
    P("Mukesh Patel School of Technology Management & Engineering", b=True, al="c", sz=12, after=0, line=1.0)


def acknowledgement():
    UNNUM_TITLE("ACKNOWLEDGEMENT", "ack", size=16, front=True, underline=True)
    J("I would like to express my special thanks and gratitude to my guide, Dr. Suresh Kurumbanshi, "
      "Head of Department, Department of Computer Engineering, SVKM’S NMIMS, Mukesh Patel School of "
      "Technology Management & Engineering, Shirpur Campus (M.H.). His constant encouragement kept me "
      "working steadily, and his close reading of my claims pushed me to check them properly; even a "
      "citation-numbering mistake in an early draft of my literature review did not get past him. He "
      "made it clear to me that a system that runs is not the same thing as a system that has been "
      "measured. I am highly indebted to him for his guidance, for the resources he provided, and for "
      "his support in completing this capstone project on time.")
    J("I also thank Dr. Venkatadri Marriboyina, Associate Dean, SVKM’S NMIMS, Mukesh Patel School of "
      "Technology Management & Engineering, Shirpur Campus (M.H.), for his timely support and the "
      "encouraging environment he created. A special mention of thanks goes to the Capstone Project "
      "Co-ordinator, Dr. Durgesh M. Sharma, Assistant Professor, Department of Computer Engineering, "
      "who took the time to hear me out and kept me on the correct path so that I could draft this "
      "report as per the specified guidelines.")
    J("I thank Kunj Kansara for helping with the relevance annotation used in the evaluation, and the "
      "invited beta testers of GitaGuide, whose use of the deployed system exposed failures in the "
      "safety gate and the answer format that I could not have found on my own. Finally, I thank all the "
      "staff members of the Computer Engineering Department and everyone who contributed, directly or "
      "indirectly, to my technical education and to this project.")
    P("Harikrishna Panchal", b=True, al="r", before=10)


def abstract():
    UNNUM_TITLE("ABSTRACT", "abstract", size=16, front=True, underline=True)
    J("Chatbots that answer in the voice of a scripture are now common. Investigative reporting in 2023 "
      "found that several widely used Bhagavad Gita chatbots told users that killing was acceptable when "
      "it fulfilled their duty, and they supported the answer with genuine verses. Retrieving a real "
      "passage and citing it correctly does not make an answer safe. This project describes GitaGuide, a "
      "guidance system built on two scriptures, the Bhagavad Gita (701 verses) and the Arthashastra "
      "(3,135 sutras), around one idea: safety has to be part of the pipeline's structure and has to be "
      "tested, rather than being hoped for from a prompt.")
    J("Every message is scored for emotion and intensity, classified into one of six intents, and passed "
      "through a four-layer crisis gate that blocks when a check cannot be completed before any retrieval or generation takes place. "
      "Guidance is retrieved from a 3,836-verse knowledge graph. Every verse identifier in the answer is "
      "checked against the retrieved set, and every stage of every turn is stored so that it can be replayed.")
    J("The system was measured end to end. The offline safety suite blocked all 36 crisis prompts and none "
      "of 25 benign controls, which shows that earlier fixes still work and is not an estimate of field performance. A "
      "conversational router removed retrieval from all 20 non-guidance turns, and INT8 quantisation kept "
      "peak memory at 429 MiB on a 512 MiB instance. Retrieval is weak (hit rate at five of 0.175 on 40 "
      "queries), and ablations attribute the measured accuracy to the dense encoder and score fusion "
      "rather than to the graph. The emotion model reaches a held-out macro-F1 of 0.451, live intent "
      "accuracy is 0.650, and a preliminary rating of 30 answers gave faithfulness of only 2.77 out of 5.")
    J("A red-team review of 33 live replies found two answers that presented scripture-sanctioned violence "
      "as permitted. Both were traced to a gap in the output gate. A fifth output rule was added, flagged "
      "both replies and none of the other 31 when replayed, and has since been deployed to production. Its "
      "performance on unseen prompts has not been measured. The contribution of the project is a pipeline in "
      "which safety, citation checking and self-measurement are built in, with an honest account of where "
      "it still falls short.")


def _toc_row(t, num, title, pg, bold, indent, size=11):
    cells = t.add_row().cells
    for c, txt, al in ((cells[0], num, None), (cells[1], ("    " * indent) + title, None), (cells[2], pg, "r")):
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(1)
        if al == "r":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(txt)
        r.bold = bold
        r.font.size = Pt(size)
    for c, w in zip(cells, (1.6, 11.0, 2.0)):
        c.width = Cm(w)


def toc():
    reg = old_registry()
    UNNUM_TITLE("TABLE OF CONTENTS", "toc", size=14, front=True)
    t = D.add_table(rows=1, cols=3)
    t.autofit = False
    for c, txt in zip(t.rows[0].cells, ("Sr. No.", "Chapter / Section", "Page")):
        c.text = ""
        r = c.paragraphs[0].add_run(txt)
        r.bold = True
        r.font.size = Pt(11)
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
    t.rows[0].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for c, w in zip(t.rows[0].cells, (1.6, 11.0, 2.0)):
        c.width = Cm(w)
    _toc_row(t, "", "List of Figures", page_of("f_lof"), True, 0)
    _toc_row(t, "", "List of Tables", page_of("f_lot"), True, 0)
    _toc_row(t, "", "List of Abbreviations", page_of("f_abbr"), True, 0)
    for kind, bm, label, title in reg:
        if kind == "ch":
            _toc_row(t, label, title.upper(), page_of(bm), True, 0)
        elif kind == "h2":
            _toc_row(t, label, title, page_of(bm), False, 1)
        elif kind == "h3":
            _toc_row(t, label, title, page_of(bm), False, 2, size=10.5)
        elif kind == "refs":
            _toc_row(t, "", "References", page_of(bm), True, 0)
        elif kind == "app":
            _toc_row(t, "", f"Appendix {label}: {title}", page_of(bm), True, 0)


def lof_lot():
    reg = old_registry()
    UNNUM_TITLE("LIST OF FIGURES", "lof", size=14, front=True)
    rows = [[i, f"Figure {lab}", title, page_of(bm)] for i, (k, bm, lab, title) in
            enumerate([r for r in reg if r[0] == "fig"], 1)]
    _plain_list(["Sr. No.", "Figure No.", "Name of Figure", "Page"], rows, [1.8, 2.3, 8.9, 1.6])
    UNNUM_TITLE("LIST OF TABLES", "lot", size=14, front=True)
    rows = [[i, f"Table {lab}", title, page_of(bm)] for i, (k, bm, lab, title) in
            enumerate([r for r in reg if r[0] == "tbl"], 1)]
    _plain_list(["Sr. No.", "Table No.", "Name of Table", "Page"], rows, [1.8, 2.3, 8.9, 1.6])


def _plain_list(head, rows, widths):
    t = D.add_table(rows=1, cols=len(head))
    t.autofit = False
    for c, h, w in zip(t.rows[0].cells, head, widths):
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(11)
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
        c.width = Cm(w)
    for row in rows:
        cells = t.add_row().cells
        for k, (c, v, w) in enumerate(zip(cells, row, widths)):
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            if k in (0, 1, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if k != 3 else WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(str(v))
            r.font.size = Pt(10.5)
            c.width = Cm(w)


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"),
    ("ARSH", "Abrupt Refusal Secondary Harm"), ("BG", "Bhagavad Gita (verse identifier prefix)"),
    ("AS", "Arthashastra (sutra identifier prefix)"), ("BM25", "Best Matching 25, a lexical ranking function"),
    ("CI", "Continuous Integration"), ("CPU", "Central Processing Unit"),
    ("EC2", "Elastic Compute Cloud (Amazon Web Services)"), ("GoEmotions", "A 28-label emotion dataset from Google Research"),
    ("HTTP", "Hypertext Transfer Protocol"), ("INT8", "8-bit integer quantisation"),
    ("JSON", "JavaScript Object Notation"), ("KG", "Knowledge Graph"),
    ("LLM", "Large Language Model"), ("MiB", "Mebibyte (1,048,576 bytes)"),
    ("MRR", "Mean Reciprocal Rank"), ("NIM", "NVIDIA Inference Microservices"),
    ("nDCG", "Normalised Discounted Cumulative Gain"), ("ONNX", "Open Neural Network Exchange"),
    ("RAG", "Retrieval-Augmented Generation"), ("RoBERTa", "Robustly Optimised BERT Pretraining Approach"),
    ("SQL", "Structured Query Language"), ("TLS", "Transport Layer Security"),
    ("UI", "User Interface"), ("vCPU", "Virtual CPU"),
]


def abbreviations():
    UNNUM_TITLE("LIST OF ABBREVIATIONS", "abbr", size=14, front=True)
    t = D.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    for c, h, w in zip(t.rows[0].cells, ("Abbreviation", "Expansion"), (3.5, 11.0)):
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(11)
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
        c.width = Cm(w)
        lib._shade(c, "E7ECF4")
    for a, b in ABBR:
        cells = t.add_row().cells
        for c, v, w in zip(cells, (a, b), (3.5, 11.0)):
            c.text = ""
            c.paragraphs[0].paragraph_format.line_spacing = 1.15
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            r = c.paragraphs[0].add_run(v)
            r.font.size = Pt(11)
            c.width = Cm(w)
