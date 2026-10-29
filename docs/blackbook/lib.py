# -*- coding: utf-8 -*-
"""Helpers for the NMIMS Black Book build (template: Capstone Project Report Format)."""
import json
import os
import re

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PAGES_FILE = os.path.join(HERE, "_pages.json")
NUMS_FILE = os.path.join(HERE, "_numbers.json")

PROJECT_TITLE = "GitaGuide: Emotion-Aware, Explainable, Safety-Gated Guidance"
YEAR = "2025-26"
TEXT_W = 21.0 - 3.81 - 2.54  # cm, A4 minus 1.5" left and 1" right margins

PAGES = json.load(open(PAGES_FILE, encoding="utf-8-sig")) if os.path.exists(PAGES_FILE) else {}
OLD_NUMS = json.load(open(NUMS_FILE, encoding="utf-8")) if os.path.exists(NUMS_FILE) else {}
NUMS = {}
REG = []  # (kind, bookmark, label, title)
_state = {"chapter": 0, "tbl": 0, "fig": 0, "bm": 100, "appendix": None}

D = docx.Document()


def roman(n):
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"),
            (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def page_of(bm):
    raw = PAGES.get(bm)
    if raw is None:
        return "?"
    return roman(raw) if bm.startswith("f_") else str(raw)


def _set_fonts(style, name="Times New Roman"):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + a), name)


def setup():
    sec = D.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.81), Cm(2.54)
    sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)
    sec.header_distance = sec.footer_distance = Cm(1.25)

    n = D.styles["Normal"]
    n.font.name = "Times New Roman"
    n.font.size = Pt(12)
    _set_fonts(n)
    n.paragraph_format.line_spacing = 1.5
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.space_before = Pt(0)

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        s = D.styles[name]
        s.font.name = "Times New Roman"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
        _set_fonts(s)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.line_spacing = 1.5
        s.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 0)
        s.paragraph_format.space_after = Pt(6)
    D.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    D.styles["Heading 1"].paragraph_format.space_after = Pt(18)
    for name in ("List Bullet", "List Number"):
        s = D.styles[name]
        s.font.name = "Times New Roman"
        s.font.size = Pt(12)
        _set_fonts(s)
        s.paragraph_format.line_spacing = 1.5
        s.paragraph_format.space_after = Pt(3)


# ------------------------------------------------------------------ bookmarks
def bookmark(p, name):
    _state["bm"] += 1
    bid = str(_state["bm"])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(start)
    else:
        p._p.insert(0, start)
    p._p.append(end)


# ------------------------------------------------------------------ text
def _runs(p, text, size=None, bold=None, italic=None, font=None):
    """Add text to p; `_{x}` becomes subscript and `^{x}` superscript."""
    parts = re.split(r"(_\{[^}]*\}|\^\{[^}]*\})", text)
    for part in parts:
        if not part:
            continue
        sub = part.startswith("_{")
        sup = part.startswith("^{")
        r = p.add_run(part[2:-1] if (sub or sup) else part)
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic
        if font:
            r.font.name = font
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    return p


def P(text="", b=None, i=None, al=None, sz=None, after=None, before=None, keep=False, u=False, line=None):
    p = D.add_paragraph()
    pf = p.paragraph_format
    if after is not None:
        pf.space_after = Pt(after)
    if before is not None:
        pf.space_before = Pt(before)
    if line is not None:
        pf.line_spacing = line
    if al == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif al == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif al == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if keep:
        pf.keep_with_next = True
    if text:
        _runs(p, text, size=sz, bold=b, italic=i)
        if u:
            for r in p.runs:
                r.underline = True
    return p


def J(text):
    return P(text, al="j")


def LEAD(label, text):
    """Paragraph that opens with a bold run-in label."""
    p = D.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(label + " ")
    r.bold = True
    _runs(p, text)
    return p


def B(label, text=None):
    p = D.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text is None:
        _runs(p, label)
    else:
        r = p.add_run(label + " ")
        r.bold = True
        _runs(p, text)
    return p


def NUM(n, label, text=None):
    p = D.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.space_after = Pt(3)
    if text is None:
        p.add_run(f"{n}.\t")
        _runs(p, label)
    else:
        p.add_run(f"{n}.\t")
        r = p.add_run(label + " ")
        r.bold = True
        _runs(p, text)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    return p


def SUMMARY(text):
    return LEAD("Chapter summary.", text)


# ------------------------------------------------------------------ headings
def CH(num, title, key):
    _state["chapter"] = num
    _state["tbl"] = 0
    _state["fig"] = 0
    _state["appendix"] = None
    p = D.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    r = p.add_run(f"Chapter {num}")
    r.add_break(WD_BREAK.LINE)
    p.add_run(title)
    bookmark(p, f"bm_{key}")
    REG.append(("ch", f"bm_{key}", str(num), title))
    return p


def H2(num, title, key):
    p = D.add_paragraph(style="Heading 2")
    p.add_run(f"{num} {title}")
    bookmark(p, f"bm_{key}")
    REG.append(("h2", f"bm_{key}", num, title))
    return p


def H3(num, title, key=None):
    p = D.add_paragraph(style="Heading 3")
    p.add_run(f"{num} {title}")
    if key:
        bookmark(p, f"bm_{key}")
        REG.append(("h3", f"bm_{key}", num, title))
    return p


def UNNUM_TITLE(title, key, size=16, front=False, page_break=True, align="c", underline=False):
    """Un-numbered front/back matter title (Certificate, References, Appendix ...)."""
    p = D.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(title)
    r.font.size = Pt(size)
    r.underline = underline
    bookmark(p, f"{'f_' if front else 'bm_'}{key}")
    return p


# ------------------------------------------------------------------ numbering for tables and figures
def _label(kind, key):
    ap = _state["appendix"]
    if ap:
        _state[kind] += 1
        return f"{ap}.{_state[kind]}"
    _state[kind] += 1
    return f"{_state['chapter']}.{_state[kind]}"


def T(key):
    return "Table " + OLD_NUMS.get("tbl:" + key, "?")


def F(key):
    return "Figure " + OLD_NUMS.get("fig:" + key, "?")


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), color)
    tcPr.append(s)


def _cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def _fmt_cell(cell, text, size, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, str(text), size=size, bold=bold)


def TBL(key, caption, head, rows, widths=None, size=10, center_cols=()):
    """Table with the caption ABOVE it (template rule)."""
    label = _label("tbl", key)
    NUMS["tbl:" + key] = label
    cap = P(f"Table {label}: {caption}", al="c", sz=10, after=3, keep=True, line=1.0, b=True)
    bookmark(cap, f"bm_tbl_{key}")
    REG.append(("tbl", f"bm_tbl_{key}", label, caption))
    t = D.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _cell_margins(t)
    for k, h in enumerate(head):
        c = t.rows[0].cells[k]
        _fmt_cell(c, h, size, bold=True, align="c")
        _shade(c, "E7ECF4")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for row in rows:
        cs = t.add_row().cells
        for k, v in enumerate(row):
            _fmt_cell(cs[k], v, size, align="c" if k in center_cols else None)
    if widths:
        tot = sum(widths)
        scale = TEXT_W / tot if tot > TEXT_W else 1.0
        for row in t.rows:
            for k, w in enumerate(widths):
                row.cells[k].width = Cm(w * scale)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    P("", after=4, line=1.0)
    return t


def FIG(key, path, caption, w=12.0):
    """Figure with the caption BELOW it (template rule)."""
    label = _label("fig", key)
    NUMS["fig:" + key] = label
    p = D.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(os.path.join(HERE, path), width=Cm(min(w, TEXT_W)))
    cap = P(f"Figure {label}: {caption}", al="c", sz=10, after=10, line=1.0, b=True)
    bookmark(cap, f"bm_fig_{key}")
    REG.append(("fig", f"bm_fig_{key}", label, caption))


def CODE(title, code, key=None):
    label = _label("tbl", "x") if False else None
    t = P(title, b=True, sz=10, after=3, keep=True, line=1.0, before=6)
    for line in code.strip("\n").split("\n"):
        p = D.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.2)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        ppr.append(shd)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
    P("", after=4, line=1.0)


def EQ(parts, number):
    p = D.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.tab_stops.add_tab_stop(Cm(TEXT_W / 2), alignment=1)  # center
    pf.tab_stops.add_tab_stop(Cm(TEXT_W), alignment=2)  # right
    p.add_run("\t")
    _runs(p, parts, italic=True)
    p.add_run(f"\t({number})")
    return p


def ALG(title, steps):
    """Boxed numbered algorithm (a one-cell table)."""
    P(title, b=True, sz=10, after=3, keep=True, line=1.0, before=4)
    t = D.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_margins(t, 80, 80, 120, 120)
    c = t.rows[0].cells[0]
    c.width = Cm(TEXT_W)
    _shade(c, "FAFAFA")
    c.text = ""
    first = True
    for i, s in enumerate(steps, 1):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        _runs(p, f"{i}.  {s}", size=10.5)
    P("", after=4, line=1.0)


def NOTE(text):
    """Highlighted placeholder / clarification flag (visible on purpose)."""
    p = P("", al="j")
    r = p.add_run(text)
    r.font.highlight_color = 7  # yellow
    return p


# ------------------------------------------------------------------ header / footer
def _field(run, instr):
    for kind, txt in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        run._r.append(e)


def header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.line_spacing = 1.0
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), alignment=2)
    r = hp.add_run(PROJECT_TITLE + "\t" + YEAR)
    r.font.size = Pt(10)
    ppr = hp._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "808080")):
        bt.set(qn("w:" + k), v)
    bdr.append(bt)
    ppr.append(bdr)
    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.line_spacing = 1.0
    fr = fp.add_run()
    fr.font.size = Pt(11)
    _field(fr, " PAGE ")


def page_numbering(section, fmt, start=1):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def new_section():
    s = D.add_section(WD_SECTION.NEW_PAGE)
    s.left_margin, s.right_margin = Cm(3.81), Cm(2.54)
    s.top_margin, s.bottom_margin = Cm(2.54), Cm(2.54)
    return s


def finish(path):
    with open(NUMS_FILE, "w", encoding="utf-8") as f:
        json.dump({k: v for k, v in NUMS.items()}, f, indent=1)
    with open(os.path.join(HERE, "_registry.json"), "w", encoding="utf-8") as f:
        json.dump(REG, f, indent=1, ensure_ascii=False)
    D.save(path)


def APPENDIX(letter, title, key):
    _state["appendix"] = letter
    _state["tbl"] = 0
    _state["fig"] = 0
    p = D.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.add_run(f"Appendix {letter}: {title}")
    bookmark(p, f"bm_{key}")
    REG.append(("app", f"bm_{key}", letter, title))
    return p


def REFHEAD(key="refs"):
    p = D.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.add_run("References")
    bookmark(p, f"bm_{key}")
    REG.append(("refs", f"bm_{key}", "", "References"))
    return p


def old_registry():
    f = os.path.join(HERE, "_registry.json")
    return json.load(open(f, encoding="utf-8")) if os.path.exists(f) else []
