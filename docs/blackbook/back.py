# -*- coding: utf-8 -*-
import os

import docx
from docx.shared import Pt

import lib
from lib import APPENDIX, B, CODE, D, H2, J, NOTE, P, REFHEAD, T, TBL

def _load_refs():
    import re
    src = docx.Document(os.path.join(lib.HERE, "..", "paper", "Research_Paper.docx"))
    out = {}
    for p in src.paragraphs:
        m = re.match(r"^\[(\d+)\]	(.+)$", p.text.strip())
        if m:
            out[int(m.group(1))] = m.group(2).strip()
    assert sorted(out) == list(range(1, 28)), sorted(out)
    return [out[i] for i in range(1, 28)]


REFS = _load_refs()


def references():
    REFHEAD()
    J("The references are numbered, and each one is cited in the text of the report.")
    for i, r in enumerate(REFS, 1):
        p = P("", al="j", sz=11, after=4, line=1.15)
        p.paragraph_format.left_indent = lib.Cm(1.0)
        p.paragraph_format.first_line_indent = lib.Cm(-1.0)
        p.add_run(f"[{i}]\t").font.size = Pt(11)
        p.add_run(r).font.size = Pt(11)
        p.paragraph_format.tab_stops.add_tab_stop(lib.Cm(1.0))


def appendix_a():
    APPENDIX("A", "Sample Code", "appA")
    J("These listings are excerpts from the project repository, shortened where marked. They show the "
      "pieces of code that carry the main design claims: the fixed order of the pipeline, the safety gate "
      "and the fifth output rule.")
    CODE("Listing A.1: The fixed pipeline order (app/pipeline/orchestrator.py, handle_turn, shortened)", '''
async def handle_turn(self, turn: UserTurn) -> GuidanceResult:
    """Run one full turn of the pipeline."""
    session = await self.repository.ensure_session(turn.session_id, turn.locale)
    message = turn.message.strip()

    # 1. emotion: CPU-bound ONNX, so off the event loop.
    emotion = await asyncio.to_thread(self.detector.detect, message)

    # 2. intent: what the message is ABOUT. Feeds the gate and the
    # scripture prior in re-ranking. Not the router.
    intent = await self.intent_classifier.classify(message, emotion)

    # 3. SAFETY GATE, and it runs here so that nothing downstream can
    # weaken it: a blocked turn returns below, before retrieval,
    # generation or the router exist in the call path at all.
    input_verdict = self.safety_gate.check_input(message, emotion, intent, session.id)
    ...  # the user message and the safety log are persisted here
    if input_verdict.blocked:
        return await self._crisis_short_circuit(...)

    # 4. memory: rolling summary + last N turns.
    memory_context = await self.memory.load(session.id)

    # 4b. router: lexicon first (0 ms, no model), else one cheap LLM call.
    # Any failure fails OPEN to `guidance`, which retrieves.
    decision = await self.router.route(message, memory_context)
    tier = tier_for(decision.route)

    # 5. retrieval, now conditional: a greeting gets no verse.
    retrieval = None
    if tier.retrieves:
        retrieval = await self.retriever.retrieve(
            decision.retrieval_query(message, memory_context), emotion, intent,
            vector_query=decision.vector_query(message, self.settings.query_rewrite))

    # 6. grounded generation.
    if retrieval is None:
        generation = await self.generator.generate_conversational(message, tier, memory_context)
    else:
        generation = await self.generator.generate(
            message, retrieval, emotion, intent, memory_context, tier)

    # 7. output gate, outside the retrieval branch: it runs on greetings too.
    output_verdict = self.safety_gate.check_output(generation.answer, session.id, input_verdict)
    answer = (self.safety_gate.crisis_response(output_verdict)
              if output_verdict.blocked else generation.answer)

    # 8. persist the full trace and refresh the rolling summary.
    ...
''')
    CODE("Listing A.2: The input and output safety checks (app/pipeline/safety.py)", '''
def check_input(self, text, emotion=None, intent=None, session_id=None) -> SafetyVerdict:
    """Gate the user's message. MUST be called before retrieval/generation."""
    triggers = self._run_rules(text, self.rules, SafetyLayer.RULE)
    rule_fired = bool(triggers)
    model_triggers, signals = self._run_risk_policy(emotion, intent)
    triggers.extend(model_triggers)
    triggers.extend(self._run_soft_cue_fallback(text, intent, rule_fired))

    verdict = SafetyVerdict(blocked=bool(triggers), triggers=triggers, signals=signals)
    if verdict.blocked:
        categories = [t.category for t in triggers]
        verdict.response = build_crisis_response(categories)

    self._log(text, verdict, emotion, intent, session_id)
    return verdict

def check_output(self, answer, session_id=None, input_verdict=None) -> SafetyVerdict:
    """Gate the generated answer (rule 4: after generation)."""
    triggers = self._run_rules(answer, self.output_rules, SafetyLayer.OUTPUT)
    verdict = SafetyVerdict(blocked=bool(triggers), triggers=triggers)
    if verdict.blocked:
        verdict.response = OUTPUT_FALLBACK_RESPONSE
        if input_verdict is not None and input_verdict.blocked:
            verdict.response = build_crisis_response(
                [t.category for t in input_verdict.triggers])
    self._log(answer, verdict, None, None, session_id, kind="output")
    return verdict
''')
    CODE("Listing A.3: The fifth output rule, added after the red-team review (app/pipeline/safety.py)", '''
_rule(
    "output_historical_violence_endorsement",
    SafetyCategory.UNSAFE_OUTPUT,
    r"\\b(?:do(?:es)? permit|permits?|permitted|sanctions?|allows?|condones?) "
    r"(?:physical correction|corporal punishment|torture|wife[- ]beating|beating (?:her|him|them|a wife))\\b"
    r"|\\btorture (?:is|was|can be|remains) (?:acceptable|justified|permissible|warranted)\\b"
    r"|\\b(?:three|a few) (?:beats?|strikes?|lashes?) (?:with|using)\\b",
    "answer operationalises a historical violence/torture method instead of refusing",
    negatable=True,   # a correct refusal ("does not permit ...") is not blocked
),
''')
def appendix_b():
    APPENDIX("B", "Data Sheets", "appB")
    J("The data sheets collect the numbers that define the corpus, the evaluation instruments and the "
      "configuration of the deployed system.")
    TBL("corpus", "Corpus and graph composition",
        ["Field", "Detail", "Role"],
        [["Bhagavad Gita verses", "701 verses, identifiers such as BG-2.47, Shri Purohit Swami translation",
          "Personal and spiritual guidance"],
         ["Arthashastra sutras", "3,135 sutras, identifiers such as AS-7.1.5, R. Shamasastry 1915 translation",
          "Professional and statecraft guidance"],
         ["Total corpus units", "3,836", "Indexed in both the vector store and the graph"],
         ["Graph relationships, before hub suppression", "281,523 of 400,000 allowed on the free tier",
          "Derived RELATED_TO edges"],
         ["Graph relationships, after suppression", "14,709 (3.7% of the budget)",
          "Labels on more than 8% of the corpus excluded"],
         ["Emotion-tag coverage", "1,253 verses (32.7%)", "Limits graph annotation by emotion"],
         ["Theme-tag coverage", "138 verses (3.6%)", "Limits graph annotation by theme"]],
        [4.4, 6.0, 4.2], size=9.5)
    TBL("params", "Main configuration parameters (app/config.py)",
        ["Parameter", "Value", "Meaning"],
        [["retrieval_candidate_k / retrieval_top_k", "20 / 5", "Candidates retrieved and verses kept"],
         ["graph_expansion_hops", "2", "Maximum hops for graph paths"],
         ["graph_path_boost, scripture_intent_boost", "0.15, 0.25", "Weights in the re-ranking formula"],
         ["kg_hub_label_max_fraction", "0.08", "Labels above this share of the corpus are hubs"],
         ["embedding_dim", "384", "all-MiniLM-L6-v2 vector size"],
         ["emotion_top_k, emotion_score_threshold", "5, 0.05", "Labels kept from the emotion model"],
         ["safety_crisis_intent_threshold", "0.5", "Semantic crisis score that blocks"],
         ["memory_last_n_turns, summary trigger, summary size", "6, 8, 1,200 characters", "Memory window"],
         ["router_timeout_seconds", "6.0", "Deadline for the router model call"],
         ["generation_budget_seconds", "45.0", "Wall-clock budget for the provider chain"],
         ["output_guard_min_words, max_repetition", "12, 0.6", "Output sanity guard limits"],
         ["llm_temperature, llm max tokens", "0.6, 1,200", "Generation settings"],
         ["pipeline_queue_timeout_seconds", "30.0", "Longest wait of a queued turn"]],
        [5.8, 3.2, 5.6], size=9.5)
def appendix_c():
    APPENDIX("C", "List of Components", "appC")
    J("This appendix lists the software components, packages and services used. Versions are those pinned "
      "in the project’s requirements file.")
    TBL("pkgs", "Python packages (requirements.txt)",
        ["Package", "Version", "Purpose"],
        [["fastapi", "0.141.1", "HTTP API framework"], ["uvicorn", "0.52.3", "ASGI server"],
         ["pydantic, pydantic-settings", "2.13.4, 2.14.2", "Typed models and configuration"],
         ["SQLAlchemy", "2.0.52", "Database access layer"], ["asyncpg", "0.31.0", "PostgreSQL async driver"],
         ["aiosqlite", "0.22.1", "SQLite driver for tests and laptop mode"],
         ["PyJWT", "2.13.0", "Verification of Clerk tokens"],
         ["onnxruntime", "1.29.0", "CPU inference for the emotion and embedding models"],
         ["transformers, tokenizers", "5.15.1, 0.22.2", "Tokenisation (torch-free serving path)"],
         ["huggingface-hub", "1.28.0", "Model download"], ["numpy", "2.4.4", "Numerical work"],
         ["neo4j, neo4j-graphrag", "6.2.0, 1.18.0", "Graph driver and hybrid retriever"],
         ["openai, anthropic", "3.1.0, 0.122.0", "Provider clients for the language-model layer"],
         ["httpx, requests", "0.28.1, 2.34.2", "HTTP clients"]],
        [4.8, 3.4, 6.4], size=9.5)
def appendix_d():
    APPENDIX("D", "List of Papers Presented and Published", "appD")
    B("Research paper.",
      "“GitaGuide: An Emotion-Aware, Explainable, Safety-Gated System for Scripture-Grounded Guidance.” "
      "Prepared for submission to a Scopus-indexed journal.")
    B("Literature review.",
      "“Emotion-Aware and Explainable AI for Scripture-Grounded Guidance: A Critical Literature Review” [12]. "
      "Prepared for submission to ICAAIC 2027.")
    NOTE("[Author to confirm before submission: the submission status of both papers and the venue name. "
         "Neither paper is stated here as presented or published.]")
